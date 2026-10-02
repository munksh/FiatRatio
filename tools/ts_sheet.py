#!/usr/bin/env python3
"""Translations as a spreadsheet, for reading through and correcting.

    python3 tools/ts_sheet.py export translations.xlsx
    python3 tools/ts_sheet.py import translations.xlsx

export  writes every text of the app, with its English source and the Swedish,
        German and Russian translations side by side, plus the standard
        categories. Needs openpyxl.
import  reads the sheet "Texter" back into translations/*.ts. Only the Swedish,
        German and Russian columns are read. Rows are matched on the columns
        "Var" (the page the text belongs to) and "Engelska" (the source), so
        change those two in the QML, not in the sheet. Run lrelease or build
        after importing.

The sheet "Kategorier" is for reading only. The names live in
tools/make_standard.py."""
import json
import os
import sys
import xml.etree.ElementTree as ET

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
LANGS = [("sv", "Svenska"), ("de", "Tyska"), ("ru", "Ryska")]


def ts_path(lang):
    return os.path.join(ROOT, "translations", "harbour-fiatratio-%s.ts" % lang)


def read_ts(lang):
    return ET.parse(ts_path(lang))


def export(out):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    trees = {l: read_ts(l) for l, _ in LANGS}
    rows = {}
    order = []
    for lang, _ in LANGS:
        for ctx in trees[lang].getroot().iter("context"):
            cname = ctx.find("name").text
            for m in ctx.iter("message"):
                tr = m.find("translation")
                if tr.get("type") in ("obsolete", "vanished"):
                    continue
                key = (cname, m.find("source").text)
                if key not in rows:
                    rows[key] = {"note": (m.find("extracomment").text if m.find("extracomment") is not None else "")}
                    order.append(key)
                rows[key][lang] = tr.text or ""

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head_fill = PatternFill("solid", fgColor="E6E1D8")
    edit_fill = PatternFill("solid", fgColor="FFF8DC")
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    ws = wb.active
    ws.title = "Texter"
    head = ["Var", "Engelska"] + [n for _, n in LANGS] + ["Anmärkning", "Din kommentar"]
    ws.append(head)
    for key in order:
        r = rows[key]
        ws.append([key[0], key[1]] + [r.get(l, "") for l, _ in LANGS] + [r["note"] or "", ""])
    widths = [16, 46, 46, 46, 46, 28, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.font = font
            c.alignment = wrap
    for c in ws[1]:
        c.font = bold
        c.fill = head_fill
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for c in row:
            c.fill = edit_fill
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    std = json.load(open(os.path.join(ROOT, "qml", "data", "standard.json"), encoding="utf-8"))
    wk = wb.create_sheet("Kategorier")
    wk.append(["Slag", "Nyckel", "Hör till", "Engelska", "Svenska", "Tyska", "Ryska"])
    for k, v in std["buckets"].items():
        wk.append(["hink", k, "", v["en"], v["sv"], v["de"], v["ru"]])
    for k, v in std["categories"].items():
        wk.append([v["kind"], k, v["parent"] or "", v["en"], v["sv"], v["de"], v["ru"]])
    for k, v in std["terms"].items():
        wk.append(["förklaring", k, "", v["en"], v["sv"], v["de"], v["ru"]])
    for i, w in enumerate([12, 22, 14, 36, 36, 36, 36], 1):
        wk.column_dimensions[get_column_letter(i)].width = w
    for row in wk.iter_rows():
        for c in row:
            c.font = font
            c.alignment = wrap
    for c in wk[1]:
        c.font = bold
        c.fill = head_fill
    wk.freeze_panes = "D2"
    wk.auto_filter.ref = wk.dimensions

    wl = wb.create_sheet("Läs mig", 0)
    lines = [
        ("Fiat Ratio, översättningar", True),
        ("", False),
        ("Fliken Texter har alla texter i appen. De gula kolumnerna (Svenska, Tyska, Ryska) är de du rättar. Rätta direkt i cellen.", False),
        ("Rör inte kolumnerna Var och Engelska. De är nyckeln som sparar tillbaka rätt text. Vill du ändra den engelska texten, säg till.", False),
        ("Skriv gärna i Din kommentar om något är oklart eller beror på sammanhanget.", False),
        ("Tecknen %1, %2 och %3 ersätts av appen med ett värde (en summa, ett datum, ett namn). De ska stå kvar, men får byta plats i meningen.", False),
        ("Fliken Kategorier är för läsning. Skriv ändringar i Din kommentar-kolumn på Texter eller i ett eget meddelande.", False),
        ("", False),
        ("Exempel på hur en rad ser ut:", True),
        ("Var: AddDialog   Engelska: Repeats   Svenska: Återkommande   Tyska: Wiederkehrend   Ryska: Повторяется", False),
    ]
    for text, is_bold in lines:
        wl.append([text])
        wl.cell(row=wl.max_row, column=1).font = bold if is_bold else font
        wl.cell(row=wl.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    wl.column_dimensions["A"].width = 120
    wb.save(out)
    print("wrote %s: %d texts" % (out, len(order)))


def import_(path):
    from openpyxl import load_workbook

    ws = load_workbook(path)["Texter"]
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    new = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key = (row[col["Var"]], row[col["Engelska"]])
        new[key] = {l: row[col[n]] for l, n in LANGS}
    for lang, _ in LANGS:
        tree = read_ts(lang)
        changed = 0
        seen = set()
        for ctx in tree.getroot().iter("context"):
            cname = ctx.find("name").text
            for m in ctx.iter("message"):
                tr = m.find("translation")
                key = (cname, m.find("source").text)
                if key not in new or tr.get("type") in ("obsolete", "vanished"):
                    continue
                seen.add(key)
                value = new[key][lang]
                if value is None or str(value).strip() == "":
                    continue
                value = str(value)
                if tr.text != value:
                    tr.text = value
                    tr.attrib.pop("type", None)
                    changed += 1
        tree.write(ts_path(lang), encoding="utf-8", xml_declaration=True)
        print("%s: %d changed" % (lang, changed))
        gone = set(new) - seen
        if gone:
            print("  rows in the sheet that match no text in the app (Var/Engelska was edited?): %d" % len(gone))


if __name__ == "__main__":
    if len(sys.argv) != 3 or sys.argv[1] not in ("export", "import"):
        print(__doc__)
        sys.exit(1)
    (export if sys.argv[1] == "export" else import_)(sys.argv[2])
